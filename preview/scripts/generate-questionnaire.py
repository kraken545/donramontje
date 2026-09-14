#!/usr/bin/env python3
"""Genera el cuestionario para el cliente en .docx y .pdf con la paleta
Don Ramon Jetfuel (negro #0B0B0B + amarillo #FFC800).

Uso: python3 generate-questionnaire.py
Salida: preview/docs/cuestionario-cliente-donramontje.docx / .pdf
"""
import os

# ---------------------------------------------------------------- contenido
SECTIONS = [
    ("A", "Identidad y marca", [
        ("1. Nombre legal del negocio (como está registrado):", "line", None),
        ('2. Nombre comercial exacto que debe aparecer en la web (¿"Don Ramon Jetfuel", "Don Ramontje", otro?):', "line", None),
        ("3. Logo oficial: ¿nos pueden enviar el archivo en alta resolución?", "checks",
         ["PNG transparente", "SVG / vector", "Solo foto del logo", "No tenemos archivo digital"]),
        ("4. ¿Cuál es el lema/tagline oficial de la marca?", "checks",
         ["Big flavor. No small portions.", "Otro: ______________________________"]),
        ("5. Idiomas de la web:", "checks",
         ["Español + Inglés", "Solo inglés", "Papiamentu", "Neerlandés", "Otro: __________"]),
    ]),
    ("B", "Contacto y pedidos", [
        ("6. WhatsApp del negocio (con código de país, ej. +5999 ...):", "line", None),
        ("7. ¿Ese WhatsApp es la vía principal para pedidos?", "checks",
         ["Sí", "No, usamos: ______________________________"]),
        ("8. Instagram (handle exacto):", "line", "@______________________________"),
        ("9. ¿Tienen Facebook, TikTok u otra red que deba aparecer?", "line", None),
        ("10. Teléfono fijo/móvil para llamadas:", "line", None),
        ("11. Correo electrónico (si aplica):", "line", None),
        ("12. Métodos de pago aceptados:", "checks",
         ["Efectivo", "Tarjeta", "Apple Pay/Google Pay", "Pago móvil local", "Otro: __________"]),
    ]),
    ("C", "Menú y precios", [
        ("13. ¿Los precios publicados en la web actual son correctos?", "checks",
         ["Sí, tal cual", "Hay cambios (escribir aquí): ______________________________"]),
        ("14. ¿Faltan platos en el menú (burgers, pasta, burritos, otros que se ven en las fotos)?", "checks",
         ["No, solo lo publicado", "Sí, agregar: ______________________________"]),
        ("15. ¿Tienen platos del día / especiales / ofertas que deban publicarse?", "checks",
         ["No", "Sí: ______________________________"]),
        ("16. ¿Qué moneda se debe mostrar primero?", "checks",
         ["XCG (NAf)", "USD", "Ambas"]),
        ("17. Información de alérgenos o avisos (ej. \"contiene gluten/lácteos\", opciones picantes):", "line", None),
    ]),
    ("D", "Ubicación y horarios", [
        ("18. ¿Dónde está el truck ahora mismo? (dirección o punto de referencia):", "line", None),
        ("19. Enlace de Google Maps del punto exacto (compartir > copiar enlace):", "line", None),
        ("20. ¿El truck es fijo o se mueve?", "checks",
         ["Fijo (misma ubicación siempre)", "Se mueve según el día", "Ambos (base fija + eventos)"]),
        ("21. Horario de atención (escribir por día):", "plain", None),
        ("    Lunes: ________   Martes: ________   Miércoles: ________   Jueves: ________", "line", None),
        ("    Viernes: ________   Sábado: ________   Domingo: ________", "line", None),
        ("22. ¿Cierran en días festivos o fechas especiales? (si aplica):", "line", None),
        ("23. ¿Cuánto tarda un pedido en estar listo (minutos aproximados)?", "line", None),
    ]),
    ("E", "Entrega y servicio", [
        ("24. ¿Hacen delivery?", "checks",
         ["No, solo pickup", "Sí, delivery propio", "Apps de delivery: __________"]),
        ("25. ¿Están en apps de delivery? (nombre de las apps):", "line", None),
        ("26. ¿Hacen catering o eventos privados (bodas, fiestas, empresas)?", "checks",
         ["No", "Sí — info para publicar: ______________________________"]),
    ]),
    ("F", "Fotos y medios", [
        ("27. ¿Nos pueden enviar las fotos en alta resolución? (las de Instagram están en baja calidad)", "checks",
         ["Sí, por WhatsApp/Drive", "No, usen las actuales"]),
        ("28. ¿Tienen fotos del truck, del equipo/staff y del grill en acción?", "checks",
         ["Sí, las enviaremos", "No tenemos"]),
        ("29. ¿Tienen vídeos o Reels que quieran destacar?", "checks",
         ["Sí: ______________________________"]),
        ("30. ¿Tienen reseñas de clientes que quieran mostrar?", "checks",
         ["No", "Sí (enviar textos o capturas): ______________________________"]),
    ]),
    ("G", "Web y dominio", [
        ("31. ¿Quieren un dominio propio (ej. donramontje.com) o sirve el gratuito (*.pages.dev)?", "checks",
         ["Dominio propio (comprado o por comprar)", "Gratuito está bien"]),
        ("32. ¿Hay algo que la web NO deba mostrar o publicar?", "line", None),
        ("33. ¿Quién mantendrá la web actualizada (horarios, ubicación, precios)? Dejamos un archivo fácil de editar:", "checks",
         ["Yo mismo/a", "Un empleado", "El desarrollador", "Otro: __________"]),
        ("34. Comentarios o peticiones adicionales:", "line", None),
    ]),
]

BLACK = "0B0B0B"
YELLOW = "FFC800"
DARK = "111111"
GRAY = "666666"

OUT_DIR = os.path.join(os.path.dirname(__file__), "..", "docs")
os.makedirs(OUT_DIR, exist_ok=True)
DOCX_PATH = os.path.join(OUT_DIR, "cuestionario-cliente-donramontje.docx")
PDF_PATH = os.path.join(OUT_DIR, "cuestionario-cliente-donramontje.pdf")
TITLE = "Información del Negocio — Don Ramon Jetfuel"

# ---------------------------------------------------------------- docx
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def bottom_border(p, color=YELLOW, sz="6"):
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:color"), color)
    bottom.set(qn("w:sz"), sz)
    bottom.set(qn("w:space"), "4")
    pbdr.append(bottom)
    pPr.append(pbdr)


def page_field(p):
    r = p.add_run()
    fld = OxmlElement("w:fldChar"); fld.set(qn("w:fldCharType"), "begin"); r._r.append(fld)
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = "PAGE"; r._r.append(instr)
    fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "end"); r._r.append(fld2)


def build_docx():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)  # A4
    for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, m, Inches(1))

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(6)
    rpr = style.element.get_or_add_rPr()
    rf = OxmlElement("w:rFonts")
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        rf.set(qn(a), "Calibri")
    rpr.append(rf)

    # título
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(48)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(TITLE)
    r.bold = True
    r.font.size = Pt(19)
    r.font.color.rgb = RGBColor.from_string(BLACK)
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(18)
    pPr = p2._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:color"), YELLOW)
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    pbdr.append(bottom)
    pPr.append(pbdr)

    intro = doc.add_paragraph()
    r = intro.add_run("Objetivo: completar la web con información 100% real para los clientes. "
                      "Marca con ☐ lo que aplique y escribe sobre las líneas. Lo que no aplique, déjalo en blanco.")
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor.from_string(GRAY)
    intro.paragraph_format.space_after = Pt(14)

    for letter, title, questions in SECTIONS:
        t = doc.add_table(rows=1, cols=1)
        t.autofit = False
        t.columns[0].width = Inches(6.27)
        cell = t.rows[0].cells[0]
        shade(cell, BLACK)
        cp = cell.paragraphs[0]
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_before = Pt(4)
        cp.paragraph_format.space_after = Pt(4)
        r = cp.add_run(f"{letter} · {title}")
        r.bold = True
        r.font.size = Pt(13)
        r.font.color.rgb = RGBColor.from_string(YELLOW)
        r.font.name = "Calibri"
        r._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

        for qtext, qtype, opts in questions:
            q = doc.add_paragraph()
            q.paragraph_format.space_after = Pt(4)
            if qtext.startswith("    "):
                q.paragraph_format.left_indent = Inches(0.4)
                qtext = qtext.strip()
            r = q.add_run(qtext)
            r.font.color.rgb = RGBColor.from_string(DARK)
            if qtype == "line":
                bottom_border(q, color=YELLOW, sz="8")
                if opts:
                    q.paragraph_format.space_after = Pt(14)
                else:
                    q.paragraph_format.space_after = Pt(10)
            elif qtype == "checks":
                for opt in opts:
                    o = doc.add_paragraph()
                    o.paragraph_format.left_indent = Inches(0.45)
                    o.paragraph_format.space_after = Pt(3)
                    ro = o.add_run("☐  ")
                    ro.font.name = "Segoe UI Symbol"
                    ro2 = o.add_run(opt)
                    ro2.font.color.rgb = RGBColor.from_string(DARK)

    # pie de página
    footer = sec.footer
    fp = footer.paragraphs[0]
    fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = fp.add_run(f"{TITLE}   |   ")
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(GRAY)
    r2 = fp.add_run("Pág. ")
    r2.font.size = Pt(8)
    r2.font.color.rgb = RGBColor.from_string(GRAY)
    page_field(fp)

    doc.save(DOCX_PATH)
    print("DOCX:", DOCX_PATH)


# ---------------------------------------------------------------- pdf
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import inch
from reportlab.lib.colors import HexColor
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                TableStyle, HRFlowable)
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_LEFT


def build_pdf():
    FONT = "Helvetica"
    FONT_B = "Helvetica-Bold"
    # DejaVu para ☐ si está disponible
    font_paths = [
        ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
         "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
    ]
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    for regular, bold in font_paths:
        if os.path.exists(regular):
            pdfmetrics.registerFont(TTFont("DV", regular))
            pdfmetrics.registerFont(TTFont("DVB", bold))
            FONT, FONT_B = "DV", "DVB"
            break

    black = HexColor("#" + BLACK)
    yellow = HexColor("#" + YELLOW)
    dark = HexColor("#" + DARK)
    gray = HexColor("#" + GRAY)

    body = ParagraphStyle("body", fontName=FONT, fontSize=10.5, leading=15,
                          textColor=dark, spaceAfter=5, alignment=TA_LEFT)
    title_s = ParagraphStyle("title", fontName=FONT_B, fontSize=19, leading=23,
                             textColor=black, alignment=TA_CENTER, spaceBefore=40, spaceAfter=4)
    intro_s = ParagraphStyle("intro", fontName=FONT, fontSize=9.5, leading=13,
                             textColor=gray, spaceAfter=12)
    check_s = ParagraphStyle("check", parent=body, leftIndent=16, spaceAfter=3)
    line_s = ParagraphStyle("line", parent=body, spaceAfter=12)
    line_s.borderPadding = 2

    def on_page(canvas, doc):
        canvas.saveState()
        canvas.setFont(FONT, 8)
        canvas.setFillColor(gray)
        canvas.drawString(inch, 0.55 * inch, f"{TITLE}   |")
        canvas.drawRightString(A4[0] - inch, 0.55 * inch, f"Pág. {doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(PDF_PATH, pagesize=A4,
                            leftMargin=inch, rightMargin=inch,
                            topMargin=inch, bottomMargin=0.9 * inch,
                            title="Cuestionario — Don Ramon Jetfuel",
                            author="Don Ramon Jetfuel")

    story = [Paragraph(TITLE, title_s)]
    story.append(HRFlowable(width="100%", thickness=3, color=yellow, spaceAfter=14))
    story.append(Paragraph("Objetivo: completar la web con información 100% real para los clientes. "
                           "Marca con ☐ lo que aplique y escribe sobre las líneas. Lo que no aplique, déjalo en blanco.",
                           intro_s))

    for letter, s_title, questions in SECTIONS:
        t = Table([[Paragraph(f"{letter} · {s_title}",
                              ParagraphStyle("h", fontName=FONT_B, fontSize=12.5,
                                             textColor=yellow, alignment=TA_CENTER))]],
                  colWidths=[A4[0] - 2 * inch])
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), black),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ]))
        story.append(t)
        story.append(Spacer(1, 10))

        for qtext, qtype, opts in questions:
            q = qtext.strip()
            if qtype == "line":
                story.append(Paragraph(q, body))
                story.append(HRFlowable(width="100%", thickness=1.2, color=yellow, spaceBefore=1, spaceAfter=12))
            elif qtype == "checks":
                story.append(Paragraph(q, body))
                for opt in opts:
                    story.append(Paragraph(f"☐&nbsp;&nbsp;{opt}", check_s))
            else:
                story.append(Paragraph(q, body))

    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)
    print("PDF:", PDF_PATH)


if __name__ == "__main__":
    build_docx()
    build_pdf()
    print("Listo.")
